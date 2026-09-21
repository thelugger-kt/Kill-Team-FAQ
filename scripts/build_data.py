#!/usr/bin/env python3
"""Convert the Kill Team FAQ spreadsheet into data/rulings.json for the web app.

Run this script whenever Rulings/Kill Team FAQ.xlsx is updated:

    python scripts/build_data.py

It reads every sheet in the workbook, extracts Question/Answer rows (plus
Faction/Season columns where present), and writes them out in a fixed sheet
order that matches how the site should display them (newest content first,
with new sheets requiring a manual addition to SHEET_ORDER below).
"""
import json
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SOURCE_XLSX = ROOT / "Rulings" / "Kill Team FAQ.xlsx"
OUTPUT_JSON = ROOT / "data" / "rulings.json"

# Fixed display order for known sheets. Sheets not listed here are appended
# after these, in the order they appear in the workbook, so nothing is
# silently dropped if a new sheet is added before this list is updated.
SHEET_ORDER = [
    "Core Rulings",
    "Universal Equipment",
    "Killzone - Tomb World",
    "Killzone - Volkus",
    "Season - Tomb World",
    "Season - Volkus",
    "Season - Bheta-Decima",
    "Season - Gallowdark",
]


def clean(value):
    """Normalize a cell value to a trimmed string, or None if empty."""
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def read_sheet(ws):
    """Read a worksheet into a list of record dicts based on its header row."""
    headers = []
    for cell in ws[1]:
        header = clean(cell.value)
        headers.append(header)

    records = []
    for row in ws.iter_rows(min_row=2):
        record = {}
        has_content = False
        for header, cell in zip(headers, row):
            if not header:
                continue
            value = clean(cell.value)
            record[header] = value
            if value:
                has_content = True
        if has_content and record.get("Question"):
            records.append(record)
    return records


def ordered_sheet_names(all_names):
    ordered = [name for name in SHEET_ORDER if name in all_names]
    remaining = [name for name in all_names if name not in SHEET_ORDER]
    return ordered + remaining


def main():
    if not SOURCE_XLSX.exists():
        raise SystemExit(f"Could not find source workbook: {SOURCE_XLSX}")

    wb = openpyxl.load_workbook(SOURCE_XLSX, data_only=True)
    all_names = wb.sheetnames

    sheets = []
    for name in ordered_sheet_names(all_names):
        ws = wb[name]
        entries = read_sheet(ws)
        sheets.append({"name": name, "entries": entries})

    missing = [name for name in SHEET_ORDER if name not in all_names]
    if missing:
        print(f"Warning: expected sheet(s) not found in workbook: {missing}")

    OUTPUT_JSON.parent.mkdir(parents=True, exist_ok=True)
    with OUTPUT_JSON.open("w", encoding="utf-8") as f:
        json.dump({"sheets": sheets}, f, ensure_ascii=False, indent=2)

    total_entries = sum(len(s["entries"]) for s in sheets)
    print(f"Wrote {OUTPUT_JSON} with {len(sheets)} sheets and {total_entries} entries.")


if __name__ == "__main__":
    main()
